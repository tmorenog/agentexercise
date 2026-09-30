"""Builds the Meal Squad exercise handout from the HBS case template.

Keeps the template's cover block, headers, footers, section setup and styles,
and fills the body with the site's instructions, the sample prompts (from the
repository's prompt files, as students see them) and figures rendered from the
slides and the site.
"""
import copy
import os
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

HERE = Path(__file__).parent
# The recipes repo (the site), for the prompt files. Override with RECIPES_REPO.
REPO = Path(os.environ.get('RECIPES_REPO', HERE / '..' / '..' / 'recipes')).resolve()
SITE = 'https://recipes-delta-red.vercel.app'
GROUP = '[your group name]'
TEXT_WIDTH = Inches(6.25)  # 8.5in page minus the template's 1.25in + 1in margins

doc = Document(HERE / 'template.docx')
steps = json.loads((HERE / 'steps.json').read_text())

# ---------------------------------------------------------------- styles
styles = doc.styles


def add_para_style(name, base, *, num_id=None, left=None, hanging=None, after=None, before=None, keep_next=False):
    st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = styles[base]
    pf = st.paragraph_format
    if left is not None:
        pf.left_indent = left
    if hanging is not None:
        pf.first_line_indent = -hanging
    if after is not None:
        pf.space_after = after
    if before is not None:
        pf.space_before = before
    pf.keep_with_next = keep_next
    if num_id is not None:
        pPr = st.element.get_or_add_pPr()
        numPr = OxmlElement('w:numPr')
        nid = OxmlElement('w:numId')
        nid.set(qn('w:val'), str(num_id))
        numPr.append(nid)
        pPr.insert(0, numPr)
    return st


# Bullets in the template's body font: the template's List Bullet list (numId 5),
# on a style based on "normal text no indent".
add_para_style('exercise bullet', 'normal text no indent', num_id=5, left=Inches(0.3), hanging=Inches(0.2), after=Pt(4))
add_para_style('exercise label', 'normal text no indent', after=Pt(4), before=Pt(6), keep_next=True)
prompt_style = add_para_style('prompt text', 'normal text no indent', after=Pt(0))
prompt_style.font.name = 'Consolas'
prompt_style.font.size = Pt(8.5)
prompt_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
prompt_style.paragraph_format.line_spacing = 1.0
rpr = prompt_style.element.get_or_add_rPr()
fonts = rpr.find(qn('w:rFonts'))
if fonts is None:
    fonts = OxmlElement('w:rFonts')
    rpr.insert(0, fonts)
for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
    fonts.set(qn(attr), 'Consolas')

# ---------------------------------------------------------------- body
body = doc.element.body
# The template's empty body paragraph after the cover: write from there.
last_empty = [p for p in doc.paragraphs if p.style.name == 'normal text' and not p.text.strip()][-1]
anchor = last_empty._p


def _place(p):
    """Move a just-added paragraph or table to after the anchor, in order."""
    global anchor
    el = p._p if hasattr(p, '_p') else p._tbl
    anchor.addnext(el)
    anchor = el
    return p


def para(text='', style='normal text', bold_lead=None, italic=False):
    p = doc.add_paragraph(style=style)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    add_rich(p, text, italic=italic)
    return _place(p)


def add_rich(p, text, italic=False):
    """Text with *emphasis* and `code` marks."""
    for part in re.split(r'(\*[^*]+\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('*') and part.endswith('*'):
            r = p.add_run(part[1:-1])
            r.italic = True
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9)
        else:
            r = p.add_run(part)
            r.italic = italic


def h1(text, page_break=False):
    p = doc.add_paragraph(style='h1')
    if page_break:
        p.paragraph_format.page_break_before = True
    p.add_run(text)
    return _place(p)


def h2(text):
    p = doc.add_paragraph(style='h2')
    p.add_run(text)
    return _place(p)


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style='exercise bullet')
        add_rich(p, it)
        _place(p)
    p.paragraph_format.space_after = Pt(6)


def label(text):
    p = doc.add_paragraph(style='exercise label')
    r = p.add_run(text)
    r.bold = True
    return _place(p)


FIG = [0]


def figure(image, title, source, width=TEXT_WIDTH):
    FIG[0] += 1
    t = doc.add_paragraph(style='figure title')
    t.paragraph_format.keep_with_next = True
    r = t.add_run(f'Figure {FIG[0]}')
    r.style = styles['figure title label']
    t.add_run(f'\t{title}')
    _place(t)
    pic = doc.add_paragraph(style='picture')
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.keep_with_next = True
    pic.add_run().add_picture(str(HERE / 'img' / image), width=width)
    _place(pic)
    s = doc.add_paragraph(style='source')
    r = s.add_run('Source:')
    r.style = styles['source label']
    s.add_run(f'\t{source}')
    _place(s)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def borders(table, color='B9C3BB', size=6):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(size))
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color)
        b.append(e)
    tblPr.append(b)


def cell_margins(table, top=100, bottom=100, left=140, right=140):
    tblPr = table._tbl.tblPr
    m = OxmlElement('w:tblCellMar')
    for side, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:w'), str(v))
        e.set(qn('w:type'), 'dxa')
        m.append(e)
    tblPr.append(m)


def fixed_width(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement('w:tblLayout')
    lay.set(qn('w:type'), 'fixed')
    tblPr.append(lay)
    tw = OxmlElement('w:tblW')
    tw.set(qn('w:w'), str(sum(widths)))
    tw.set(qn('w:type'), 'dxa')
    for old in tblPr.findall(qn('w:tblW')):
        tblPr.remove(old)
    tblPr.append(tw)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(w))
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            tcW = c._tc.get_or_add_tcPr().get_or_add_tcW()
            tcW.set(qn('w:w'), str(w))
            tcW.set(qn('w:type'), 'dxa')


def prompt_box(path):
    text = (REPO / 'public' / path.lstrip('/')).read_text().strip()
    text = text.replace('{{SITE}}', SITE).replace('{{GROUP}}', GROUP)
    lab = label('Sample prompt')
    lab.add_run('  (copy it into Lovable as it is, or change it)').italic = True
    tbl = doc.add_table(rows=1, cols=1)
    fixed_width(tbl, [9000])
    borders(tbl)
    cell_margins(tbl)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    shade(cell, 'F3F5F1')
    first = cell.paragraphs[0]
    lines = text.split('\n')
    for i, line in enumerate(lines):
        p = first if i == 0 else cell.add_paragraph()
        p.style = styles['prompt text']
        if not line.strip():
            p.paragraph_format.space_after = Pt(0)
            p.add_run('')
            continue
        m = re.match(r'^(\s*)([-•*])\s+(.*)$', line)
        if m:
            p.paragraph_format.left_indent = Inches(0.18 + 0.18 * (len(m.group(1)) // 2))
            p.paragraph_format.first_line_indent = -Inches(0.14)
            p.add_run('- ' + m.group(3))
        else:
            p.add_run(line)
    # Rows may break across pages; the label stays with the box.
    _place(tbl)
    spacer = doc.add_paragraph(style='normal text no indent')
    spacer.paragraph_format.space_after = Pt(2)
    _place(spacer)


def step_block(s, *, show_goal=True):
    h2(f"{s['num'].replace(' · warm-up', ' (warm-up)')}: {s['title']}")
    if show_goal and s['goal']:
        para(s['goal'], style='normal text no indent', bold_lead='Goal. ')
    for c in s['callouts']:
        para(c, style='normal text no indent', italic=True)
    if s['prompt']:
        prompt_box(s['prompt'])
    if s['see']:
        label('What you should see')
        bullets(s['see'])
    if s['tryit']:
        label('Things you can try')
        bullets(s['tryit'])
    for n in s['notes']:
        para(n, style='normal text no indent', italic=True)


# ---------------------------------------------------------------- content
para('Today you’ll work on an agent system that creates a varied and affordable five-day dinner plan. '
     'In groups, you’ll build AI agents in Lovable and connect them into a system in which each agent does one job and hands its work on to the next.')
para('The exercise lets you learn by building. Along the way you’ll see how agents:')
bullets([
    'receive instructions and goals, and decide for themselves how to reach them;',
    'use tools and real data, rather than making things up;',
    'follow rules they read from a shared server using MCP, instead of having them written into their code;',
    'hand their work to other agents through a coordinator;',
    'deal with rejections, surprises and failures, and explain them honestly.',
])
para('The goal is understanding, not a perfect system. Agent responses may be inconsistent, handoffs may not work as expected, and some steps may fail. '
     'That’s part of the process: pay attention to what works, investigate what doesn’t, and use each result as an opportunity to improve your system. '
     f'Everything you need is also on the exercise site, {SITE}, where the sample prompts come ready to copy with your group name filled in.')

h1('The Challenge')
para('The task is simple to state and hard to do well: five dinners, Monday to Friday, that are varied and within a budget. '
     'The recipes are real, from TheMealDB, and the prices are real Kroger prices. No single agent does it all: each one does one job and hands its work on.')
figure('crop-challenge.png', 'Today’s challenge', 'Meal Squad class slides.')

h1('The Agent System')
para('The system has four agents and a coordinator that connects them:')
bullets([
    '*Recipe Scout* (you build it): researches and proposes recipes.',
    '*Recipe Pricer* (already built): calculates the cost of each recipe.',
    '*Meal Planner* (you build it): creates a five-day plan using the priced recipes.',
    '*Shopper* (already built): chooses the class’s best plan and builds its Kroger cart.',
])
para('Every group builds a Recipe Scout; groups that have time go on to build a Meal Planner.')
figure('crop-system.png', 'Four agents and the coordinator that connects them', 'Meal Squad class slides.')

h2('What Is an Agent?')
bullets([
    '*An action* (also called a tool) performs one specific task: search for a recipe, look up its ingredients, save it. It is ordinary code that always does the same thing.',
    '*An agent* is an AI model with a goal, a set of tools and a loop: it decides which actions to use and in what order, looks at each result, and decides what to do next until the job is done, or stops and explains why it can’t finish.',
])
figure('crop-agent.png', 'An action and an agent', 'Meal Squad class slides.')

h2('The Coordinator')
para('The coordinator is the shared hub of the system. The agents never talk to each other directly: each one connects to the coordinator using MCP, learns from it which tools it offers and how to use them, and saves its results there for the next agent. The coordinator:')
bullets([
    'checks everything it receives: the format, that recipes are real TheMealDB recipes, the limits and the class’s rules for a meal plan;',
    'works out the numbers itself, so that no agent can make one up;',
    'explains any rejection, so that an agent can fix its mistake and try again;',
    'shows every request and answer, live, on the site’s Coordinator page.',
])
para('The coordinator is not an agent: it uses no AI. It is ordinary code that follows fixed rules and gives the same answer every time. That split is deliberate. '
     'Agents use judgement, which is flexible but can be wrong; the coordinator gives them firm ground, with rules that can’t be talked around. '
     'The coordinator could be an agent too: many systems use an AI “orchestrator” that decides which agent works next and reviews their work. That makes a system more flexible, but less predictable and harder to check.')
figure('crop-mcp.png', 'How an agent works with the coordinator using MCP', 'Meal Squad class slides.')

h2('Real Data')
bullets([
    '*Recipes* come from TheMealDB, a free, open database of real recipes with photos, ingredients and instructions. Your Scout Agent calls it from your Lovable app’s backend with the free developer key “1”.',
    '*Prices* come from Kroger’s API, which needs credentials, so only the agents that run on the coordinator site use it.',
    'The Recipe Pricer prices every recipe per person: the cost of one serving when the recipe is cooked for many people.',
    'At the end, the Shopper Agent scales the chosen plan to the number of people shopping (50 by default) and turns it into one Kroger cart for the week.',
    'Prices are estimates when Kroger has no match, and they are labelled as such.',
])
figure('crop-data.png', 'The two sources of real data', 'Meal Squad class slides.')

h1('Before You Start', page_break=True)
bullets([
    'Work in groups of two or three: one person builds, and everyone watches the same screen and helps.',
    'Open Lovable (lovable.dev) on the builder’s screen and log in.',
    f'Choose a unique group name, for example *team-3*, and type it on the site’s Welcome page: the site’s prompts then fill it in for you. In this document the prompts show {GROUP} instead: replace it with your group’s name.',
    'Check the class key your instructor gives you on the Welcome page. The class key goes into a Lovable secret when a prompt asks for it, never into a prompt, your page or a chat.',
])
h2('Working in Lovable')
bullets([
    'Start a new project: in Lovable, go to Projects, then New project. The Recipe Scout gets a project of its own, separate from the Meal Planner.',
    'Use Build mode. Make sure Lovable is set to Build when you paste a prompt, so that it writes the code. In Plan mode it only talks the change through.',
    'Approve Cloud. If Lovable asks to enable Cloud, approve it: the app needs a backend to reach recipes and the coordinator.',
    'Paste one step at a time and check the result before moving on.',
])
figure('crop-lovable.png', 'Starting a project in Lovable, and the Plan / Build switch', 'Meal Squad class slides; Lovable screenshot.', width=Inches(5.3))
h2('How Each Step Works')
para('Every step below has the same four parts:')
bullets([
    '*Goal:* what the step adds, and why.',
    '*Sample prompt:* to copy into Lovable, as it is or changed to make the app your own.',
    '*What you should see:* check it before moving on.',
    '*Things you can try:* where most of the learning happens. Try to break your agent, and see whether it explains itself honestly.',
])

# Part 1: the Recipe Scout
h1('Part 1: The Recipe Scout (Agent 1)', page_break=True)
para('The Recipe Scout finds real recipes that fit a dinner theme and saves them to the coordinator. You build it in stages:')
bullets([
    'a plain page (step 0);',
    'ordinary actions that search TheMealDB (step 1);',
    'the Scout Agent that decides how to use them (step 2);',
    'the connection to the coordinator, so its recipes can be priced and used by the other agents (steps 3a and 3b).',
])
figure('flow-scout.png', 'What the Scout Agent does', 'Meal Squad site, Recipe Scout page.')
scout = steps['scout']
for s in scout:
    if s['id'] == 'step-3':
        h2(f"{s['num']}: {s['title']}")
        para(s['goal'], style='normal text no indent', bold_lead='Goal. ')
        para('This step has two parts: first connect the app to the coordinator (3a), then let the Scout Agent use it (3b).', style='normal text no indent')
        continue
    step_block(s)
para('Once your Scout Agent works, see what the Pricer Agent does with your recipes on the site’s Recipe Pricer page, and then build the Meal Planner.', italic=True)

# Part 2: the Recipe Pricer
h1('Part 2: The Recipe Pricer (Agent 2, Already Built)')
para('The Pricer Agent turns every recipe the Scout Agents save into a Kroger shopping cart with prices, so the Meal Planner Agents can compare recipes by cost. Nobody has to start it: '
     'it prices each new recipe by itself, as soon as a Scout Agent saves one, within a few minutes. For each recipe it:')
bullets([
    'reads the recipe from the coordinator;',
    'searches Kroger for every ingredient;',
    'chooses a product or a close substitute, or estimates a price when Kroger has none;',
    'saves the cart to the coordinator, which works out the cost of one serving.',
])
figure('flow-pricer.png', 'What the Pricer Agent does', 'Meal Squad site, Recipe Pricer page.')
para('On the site’s Recipe Pricer page you can watch it work: every recipe the class saved, its cart and price, and each step the agent takes. You can filter the page by your group’s name.')

# Part 3: the Meal Planner
h1('Part 3: The Meal Planner (Agent 3)', page_break=True)
para('If there’s time, build the Meal Planner in a new Lovable project. Its agent:')
bullets([
    'reads the class’s priced recipes from the coordinator;',
    'drafts five dinners;',
    'asks the coordinator to check them against the budget and the class’s rules;',
    'fixes what fails, and saves the week for the Shopper Agent.',
])
figure('flow-planner.png', 'What the Meal Planner Agent does', 'Meal Squad site, Meal Planner page.')
for s in steps['planner']:
    step_block(s)
para('Once your Meal Planner saves a plan, see the Shopper Agent on the site’s Shopper page, and every plan the class saved, with each check the coordinator ran, on the Coordinator page.', italic=True)

# Part 4: the Shopper
h1('Part 4: The Shopper (Agent 4, Already Built)')
para('The Shopper Agent does the class’s shopping. It:')
bullets([
    'compares the meal plans the Meal Planner Agents saved, chooses the one the class will cook, and credits the group that made it;',
    'builds the week’s Kroger cart for it, for 50 people by default;',
    'checks each product at Kroger that day, buys an ingredient that several dinners share only once, and replaces what the store no longer carries.',
])
para('It runs on the site: the instructor starts it, or it starts by itself once enough plans are saved. The class’s plan, the cart and why the Shopper Agent chose it appear on the site’s Shopper page.')
figure('flow-shopper.png', 'What the Shopper Agent does', 'Meal Squad site, Shopper page.')

# Troubleshooting
h1('When Something Goes Wrong')
para('A step that fails is part of the exercise. You don’t need to start over:')
bullets([
    'Tell Lovable what you expected and what happened instead, and paste any error message you see.',
    'Look at what the agent did (its trace) before changing anything.',
    'If a change made things worse, ask Lovable to go back to the previous version.',
])
rows = [
    ('“Missing or wrong class key” (error 401)', 'Re-enter the CLASS_KEY secret in Lovable, then ask: “Check that every request to the coordinator sends the header Authorization: Bearer followed by the CLASS_KEY secret.”'),
    ('An error about “X-Group” (error 400)', 'The coordinator needs your group name on every request. Ask Lovable: “Send our group name in the header X-Group on every request to the coordinator.”'),
    ('The coordinator rejected a recipe or a plan', 'That’s the coordinator doing its job, and the rejection lists every problem. A good agent reads the reasons, fixes them and tries again. If it keeps making the same mistake, paste the reason into Lovable and ask it to make sure the agent passes the rejection back to the AI.'),
    ('A “CORS” error in the browser', 'Your page is calling the coordinator directly from the browser. Ask Lovable: “Move every call to the coordinator into a backend function.”'),
    ('The agent stops early or goes round in circles', 'Look at the last few steps it took. Tell Lovable what it did and what it should have done instead, or make its instructions clearer.'),
    ('Where can I see what my agent actually did?', 'On the site’s Coordinator page: every request your agent sends and the coordinator’s answer, live, including rejections and why. Filter by your group name.'),
]
tbl = doc.add_table(rows=len(rows) + 1, cols=2)
fixed_width(tbl, [2800, 6200])
borders(tbl, color='BFBFBF', size=4)
cell_margins(tbl, top=60, bottom=60, left=100, right=100)
for i, (a, b) in enumerate([('Problem', 'What to do')] + rows):
    for j, t in enumerate((a, b)):
        c = tbl.rows[i].cells[j]
        p = c.paragraphs[0]
        p.style = styles['th fl' if i == 0 else 'tb fl']
        if i:
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
        p.add_run(t)
        if i == 0:
            shade(c, 'E6E9E4')
tbl.rows[0].repeat_header = True
for row in tbl.rows:
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
tbl.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
_place(tbl)
para('', style='normal text no indent')

# Glossary
h1('Glossary')
GLOSSARY = [
    ('Action (tool)', 'Ordinary code that performs one specific task, such as searching for a recipe or saving it. It does the same thing every time. AI models call these “tools”.'),
    ('Agent', 'An AI model with a goal, a set of tools and a loop: it decides which tools to use and in what order, looks at each result, and decides what to do next until the job is done or it stops and explains why.'),
    ('Backend', 'The part of an app that runs on a server rather than in the browser. Secrets such as the class key, and every call to the coordinator, belong there.'),
    ('Class key', 'The password your instructor gives you on the Welcome page. The coordinator only answers requests that carry it. Keep it in a Lovable secret, never in a prompt, page or chat.'),
    ('Contract', 'What the coordinator tells each agent through its get_contract tool: the formats it accepts, the rules it checks and the limits that apply.'),
    ('Coordinator', 'The shared hub of the system. Ordinary code, with no AI, that checks every request, saves the results and hands them to the next agent. The agents never talk to each other directly.'),
    ('CORS error', 'A browser error that appears when a page calls another site directly. Moving the call into a backend function fixes it.'),
    ('Group name', 'Your group’s unique name (for example team-3). It is sent with every request so the coordinator knows whose work it is.'),
    ('Handoff', 'The moment one agent’s saved work becomes the next agent’s input, through the coordinator.'),
    ('Kroger', 'A US grocery chain whose API lists the products in a store, with today’s prices and sizes. The Pricer and Shopper Agents use it.'),
    ('Lovable', 'The tool you use to build your apps by describing them in plain language. Build mode writes the code; Plan mode only talks the change through.'),
    ('Lovable Cloud', 'The backend Lovable adds to your app when you approve it, so the app can call TheMealDB and the coordinator from a server.'),
    ('Lovable secret', 'A protected value stored in your Lovable project, such as the class key. The app can use it, but it never appears in your code or page.'),
    ('MCP (Model Context Protocol)', 'A standard way for an app to discover what another system can do. Your app asks the coordinator which tools it offers and gets each tool’s name, description and inputs.'),
    ('Meal plan', 'Five dinners, Monday to Friday, saved by a Meal Planner Agent with a budget and a reason for each choice.'),
    ('Popularity', 'How many groups scouted the same recipe. The Meal Planner uses it to break ties.'),
    ('Prompt', 'The instructions you type into Lovable. The sample prompts in this document can be pasted as they are or changed.'),
    ('Rejection', 'The coordinator’s answer when a save breaks a rule. It lists every problem, so a good agent reads it, fixes the mistake and tries again.'),
    ('Rules (checks)', 'The conditions the coordinator checks for every meal plan, such as the budget or the number of different cuisines. Your instructor chooses them.'),
    ('TheMealDB', 'A free, open database of real recipes with photos, ingredients and instructions, used by the Scout Agents.'),
    ('Trace', 'The step-by-step record of what an agent did: each tool it called and the answer it got. Shown in each app’s “What the agent is doing” panel and on the site’s Coordinator page.'),
]
tbl = doc.add_table(rows=len(GLOSSARY) + 1, cols=2)
fixed_width(tbl, [2400, 6600])
borders(tbl, color='BFBFBF', size=4)
cell_margins(tbl, top=60, bottom=60, left=100, right=100)
for i, (a, b) in enumerate([('Term', 'Meaning')] + GLOSSARY):
    for j, t in enumerate((a, b)):
        c = tbl.rows[i].cells[j]
        p = c.paragraphs[0]
        p.style = styles['th fl' if i == 0 else 'tb fl']
        if i:
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t)
        if i and j == 0:
            r.bold = True
        if i == 0:
            shade(c, 'E6E9E4')
for row in tbl.rows:
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
tbl.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
_place(tbl)
para('', style='normal text no indent')

# The template's own empty body paragraph goes.
last_empty._p.getparent().remove(last_empty._p)

out = HERE.parent / 'Meal Squad Exercise.docx'
doc.save(out)
print('saved', out, 'figures:', FIG[0])
